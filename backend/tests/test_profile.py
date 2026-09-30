"""Steckbrief: Luecken zwischen Plan, Stueckliste, Klemmenplan, AWL und Symboltabelle."""

import csv
from pathlib import Path

import openpyxl
import pypdfium2 as pdfium

from app.ingestion import awl_parser
from app.ingestion.pdf_layout import known_sheets
from app.ingestion.profile import CORE_DOC_TYPES, GAP_LABELS, Occurrence, build_profile
from app.ingestion.tags import extract_tags

EXAMPLE = Path(__file__).resolve().parents[2] / "examples" / "foerderband"
ALL = set(CORE_DOC_TYPES)


def occ(tag: str, doc_type: str, tag_type: str | None = None) -> Occurrence:
    if tag_type is None:
        tag_type = "terminal" if tag.startswith("-X") else "device" if tag.startswith("-") else "cross_ref" if tag.startswith("/") else "plc_address"
    return Occurrence(tag, tag_type, doc_type)


def kinds(profile: dict) -> list[tuple[str, str]]:
    return [(g["kind"], g["tag"]) for g in profile["gaps"]]


def test_device_gaps_between_schematic_and_bom():
    profile = build_profile(
        [occ("-K1", "schematic"), occ("-K1", "bom"), occ("-K2", "schematic"), occ("-Q9", "bom"), occ("-W1", "schematic")],
        {"schematic", "bom"},
    )
    assert kinds(profile) == [("device_not_in_bom", "-K2"), ("device_not_in_schematic", "-Q9")]
    assert profile["summary"] == {"devices": 4, "terminals": 0, "plc_addresses": 0, "gaps": 2}


def test_rules_only_apply_when_both_doc_types_exist():
    profile = build_profile([occ("-K2", "schematic")], {"schematic"})
    assert profile["gaps"] == []


def test_terminal_gaps_ignore_punctuation_artifacts():
    profile = build_profile(
        [occ("-X3:1", "schematic"), occ("-X3:1", "terminal_plan"), occ("-X3:2", "schematic"), occ("-X3:6.", "schematic"),
         occ("-X3:9/-X3", "schematic"), occ("-X3:7", "terminal_plan"), occ("-X3", "schematic")],
        {"schematic", "terminal_plan"},
    )
    assert kinds(profile) == [("terminal_not_in_plan", "-X3:2"), ("terminal_not_in_schematic", "-X3:7")]
    assert profile["summary"]["terminals"] == 3


def test_plc_gaps_and_io_card_tags_are_not_addresses():
    profile = build_profile(
        [occ("E0.0", "plc_program"), occ("E0.0", "plc_symbols"), occ("E0.0", "schematic"),
         occ("A4.0", "plc_program"), occ("M10.1", "plc_symbols"), occ("M5.0", "plc_program"), occ("M5.0", "plc_symbols"),
         occ("A1.1", "plc_program"), occ("-A1.1", "plc_program")],
        {"schematic", "plc_program", "plc_symbols"},
    )
    assert kinds(profile) == [
        ("address_without_symbol", "A4.0"),
        ("symbol_unused", "M10.1"),
        ("address_not_in_schematic", "A4.0"),
    ]


def test_manual_only_devices_and_dangling_cross_refs():
    profile = build_profile(
        [occ("-K1", "schematic"), occ("-K1", "bom"), occ("-K1", "manual"), occ("-K7", "manual"), occ("/3.2", "bom"), occ("/12.4", "bom")],
        {"schematic", "bom", "manual"},
        known_sheets={1, 2, 3},
    )
    assert kinds(profile) == [("device_only_in_manual", "-K7"), ("cross_ref_dangling", "/12.4")]
    assert all(g["message"] == GAP_LABELS[g["kind"]] for g in profile["gaps"])


def test_coverage_sorted_naturally_and_full_tags_collapsed():
    profile = build_profile(
        [occ("-K12", "schematic"), occ("=A1+S1-K12", "schematic"), occ("-K2", "bom"), occ("-X1:10", "schematic"), occ("-X1:2", "schematic")],
        set(),
    )
    assert [c["tag"] for c in profile["coverage"]] == ["-K2", "-K12", "-X1:2", "-X1:10"]
    assert profile["coverage"][1]["docs"] == {"schematic": 1}


def _example_occurrences() -> list[Occurrence]:
    """Kennzeichen der FB-01-Dateien wie die Pipeline sie sieht (PDF-Text per pdfium statt Docling)."""
    found: list[Occurrence] = []

    def add(text: str, doc_type: str, loose: bool = False) -> None:
        found.extend(Occurrence(t.tag, str(t.tag_type), doc_type) for t in extract_tags(text, plc_loose=loose))

    pdf = pdfium.PdfDocument(str(EXAMPLE / "01_Stromlaufplan_FB-01.pdf"))
    try:
        for index in range(len(pdf)):
            add(pdf[index].get_textpage().get_text_range(), "schematic")
    finally:
        pdf.close()
    for row in openpyxl.load_workbook(EXAMPLE / "02_Stueckliste_FB-01.xlsx").active.iter_rows(values_only=True):
        add(" | ".join(str(c) for c in row if c is not None), "bom")
    with (EXAMPLE / "03_Klemmenplan_FB-01.csv").open(encoding="utf-8-sig") as handle:
        for row in csv.reader(handle, delimiter=";"):
            add(" | ".join(row), "terminal_plan")
    for block in awl_parser.parse_awl(awl_parser.read_text(EXAMPLE / "04_SPS_Programm_FB-01.awl")):
        add(block.declaration, "plc_program", loose=True)
        for network in block.networks:
            add(f"{network.title}\n{network.code}", "plc_program", loose=True)
    for row in awl_parser.parse_symbol_table(awl_parser.read_text(EXAMPLE / "05_Symboltabelle_FB-01.sdf")):
        add(" | ".join(row.values()), "plc_symbols", loose=True)
    add((EXAMPLE / "06_Betriebsanleitung_FB-01.md").read_text(encoding="utf-8"), "manual")
    return found


def test_example_fb01_is_consistent_except_one_unused_symbol():
    sheets = known_sheets(EXAMPLE / "01_Stromlaufplan_FB-01.pdf")
    assert sheets == set(range(1, 8))
    profile = build_profile(_example_occurrences(), ALL, sheets)
    # Stoer_Quitt (M10.1) steht in der Symboltabelle, wird im AWL aber nicht verwendet
    assert kinds(profile) == [("symbol_unused", "M10.1")]
    assert profile["summary"]["devices"] >= 20
    assert profile["summary"]["terminals"] >= 25


def test_mehrstockklemme_x2_3a_zaehlt_in_plan_und_klemmenplan():
    """-X2:3a steht auf Blatt 4 und im Klemmenplan; der Index fuehrt es als -X2:3A, das _is_pin als Anschluss nimmt."""
    profile = build_profile(_example_occurrences(), ALL, known_sheets(EXAMPLE / "01_Stromlaufplan_FB-01.pdf"))
    docs = {c["tag"]: c["docs"] for c in profile["coverage"]}
    assert {"schematic", "terminal_plan"} <= set(docs.get("-X2:3A", {}))
