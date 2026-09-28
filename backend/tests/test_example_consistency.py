"""Beispielanlage FB-01: jeder Blatt/Spalten-Verweis in Stueckliste und Klemmenplan zeigt auf die
Spalte, in der das Kennzeichen im Stromlaufplan tatsaechlich steht.

Der Chat markiert die Spalte eines Verweises (/3.8) im Plan; falsche Verweise im Beispiel
wuerden neben das Bauteil zeigen.
"""

import csv
import re
from pathlib import Path

import openpyxl
import pypdfium2 as pdfium
import pytest

from app.ingestion.pdf_layout import page_columns

EXAMPLE = Path(__file__).resolve().parents[2] / "examples" / "foerderband"
PDF = EXAMPLE / "01_Stromlaufplan_FB-01.pdf"
REF = re.compile(r"^/(\d+)\.(\d+)$")


def _references() -> list[tuple[str, str]]:
    refs = []
    sheet = openpyxl.load_workbook(EXAMPLE / "02_Stueckliste_FB-01.xlsx").active
    for row in sheet.iter_rows(values_only=True):
        tag, ref = row[0], row[5] if len(row) > 5 else None
        if isinstance(tag, str) and tag.startswith("-") and not tag.startswith("-W") and isinstance(ref, str):
            refs.append((tag, ref))
    with (EXAMPLE / "03_Klemmenplan_FB-01.csv").open(encoding="utf-8-sig") as handle:
        for row in csv.reader(handle, delimiter=";"):
            if len(row) >= 6 and row[1].startswith("-X") and REF.match(row[-1]):
                refs.append((row[1], row[-1]))
    return refs


def _columns_of(tag: str, page: int) -> list[int]:
    """Spalten, in denen das Kennzeichen auf der Seite steht (exakt, nicht als Praefix von -X3:10)."""
    columns = page_columns(PDF, page)
    pdf = pdfium.PdfDocument(str(PDF))
    try:
        pdf_page = pdf[page - 1]
        width, _ = pdf_page.get_size()
        textpage = pdf_page.get_textpage()
        text = textpage.get_text_range()
        found = []
        searcher = textpage.search(tag, match_case=True)
        while (hit := searcher.get_next()) is not None:
            start, count = hit
            following = text[start + count : start + count + 1]
            strip = re.fullmatch(r"-X\d+", tag) is not None  # Klemmleiste: ihre Klemmen -X3:1 zaehlen
            if following.isalnum() or (following in {".", ":"} and not (strip and following == ":")
                                       and text[start + count + 1 : start + count + 2].isalnum()):
                continue
            # Beschriftung steht oft links neben dem Symbol: Anfang und Ende des Kennzeichens zaehlen
            left = textpage.get_charbox(start)[0] / width
            right = textpage.get_charbox(start + count - 1)[2] / width
            found += [c.n for c in columns if c.x0 <= left <= c.x1 or c.x0 <= right <= c.x1]
        return found
    finally:
        pdf.close()


REFERENCES = _references()


def test_example_has_references():
    assert len(REFERENCES) > 30


@pytest.mark.parametrize(("tag", "ref"), REFERENCES, ids=[f"{t} {r}" for t, r in REFERENCES])
def test_reference_points_to_drawn_column(tag, ref):
    sheet, column = map(int, REF.match(ref).groups())
    assert column in _columns_of(tag, sheet), f"{tag} {ref}: im Plan in Spalten {_columns_of(tag, sheet)}"


def _bom_rows() -> list[tuple]:
    sheet = openpyxl.load_workbook(EXAMPLE / "02_Stueckliste_FB-01.xlsx").active
    return [row for row in sheet.iter_rows(values_only=True) if isinstance(row[0], str) and row[0].startswith("-")]


def test_example_meets_model_criterion_in_the_parts_list():
    """contract.md Abschnitt 5, Kriterium 2: >= 5 Einbauorte und >= 20 Teile (ohne Leitungen und Klemmleisten)."""
    rows = _bom_rows()
    devices = [row for row in rows if not row[0].startswith(("-W", "-X"))]
    locations = {row[4] for row in rows if isinstance(row[4], str) and "->" not in row[4]}
    assert len(locations) >= 5, sorted(locations)
    assert len(devices) >= 20, len(devices)


def test_example_model_from_parts_list_has_five_zones_and_twenty_parts():
    """Wie die Ingestion: Stuecklistenzeilen als Markdown-Tabelle -> Kennzeichen-Index -> Modell (build_map)."""
    from app.ingestion.machine_map import build_map
    from app.ingestion.tags import TagType, extract_tags

    sheet = openpyxl.load_workbook(EXAMPLE / "02_Stueckliste_FB-01.xlsx").active
    lines, legend = [], ""
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if value is None else str(value) for value in row]
        if cells[0].startswith("Anlage"):
            legend = cells[0]
        if cells[0].startswith("-") or cells[0] == "BMK":
            lines.append("| " + " | ".join(cells) + " |")
    table = chr(10).join(lines)  # Zeile je Stuecklistenzeile, wie Docling eine Tabelle ausgibt
    bom_rows = [(t.tag, t.context) for t in extract_tags(table) if t.tag_type == TagType.DEVICE]
    result = build_map(bom_rows, legend=legend).as_dict()
    named = [z for z in result["zones"] if z["code"] != "?"]
    assert len(named) >= 5, [z["code"] for z in result["zones"]]
    assert all(z["name"] for z in named), [(z["code"], z["name"]) for z in named]
    assert result["part_count"] >= 20 and not any(z["code"] == "?" for z in result["zones"]), result

