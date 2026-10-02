"""Testdokumentation (UR-01, PM1-AR): in sich stimmig und von allen Parsern lesbar. Ohne DB, ohne LLM.

Verweise `/Blatt.Spalte` in Stueckliste und Klemmenplan muessen auf die Spalte zeigen, in der das
Kennzeichen im Stromlaufplan steht; Stueckliste, Klemmenplan, Symboltabelle, AWL und Handbuch
muessen die vorhandenen Werkzeuge (Signalweg, Onboarding) fuettern.
"""

import csv
import re
from pathlib import Path

import openpyxl
import pypdfium2 as pdfium
import pytest

from app.api.signal import _bom_rows, _terminal_rows
from app.ingestion.awl_parser import parse_awl, parse_symbol_table, read_text
from app.ingestion.onboarding import fault_rows_from_markdown, guess_machine
from app.ingestion.pdf_layout import page_columns
from app.ingestion.signal_graph import build_graph, signal_path
from app.ingestion.tags import TagType, extract_tags

EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
SETS = {
    "UR-01": {"dir": "umroller", "start": "-S4", "safety": "-S12", "target": "-U5", "machine_type": "main"},
    "PM1-AR": {"dir": "aufrollung", "start": "-S5", "safety": "-S4", "target": "-U1", "machine_type": "main"},
}
REF = re.compile(r"^/(\d+)\.(\d+)$")


def _files(code: str) -> dict[str, Path]:
    folder = EXAMPLES / SETS[code]["dir"]
    return {
        "pdf": folder / f"01_Stromlaufplan_{code}.pdf",
        "bom": folder / f"02_Stueckliste_{code}.xlsx",
        "csv": folder / f"03_Klemmenplan_{code}.csv",
        "awl": folder / f"04_SPS_Programm_{code}.awl",
        "sdf": folder / f"05_Symboltabelle_{code}.sdf",
        "md": folder / f"06_Betriebsanleitung_{code}.md",
    }


def _references(code: str) -> list[tuple[str, str]]:
    files = _files(code)
    if not files["bom"].exists():
        return []
    refs = []
    for row in openpyxl.load_workbook(files["bom"]).active.iter_rows(values_only=True):
        tag, ref = row[0], row[5] if len(row) > 5 else None
        if isinstance(tag, str) and tag.startswith("-") and not tag.startswith("-W") and isinstance(ref, str):
            refs.append((tag, ref))
    with files["csv"].open(encoding="utf-8-sig") as handle:
        for row in csv.reader(handle, delimiter=";"):
            if len(row) >= 6 and row[1].startswith("-X") and REF.match(row[-1]):
                refs.append((row[1], row[-1]))
    return refs


def _columns_of(pdf_path: Path, tag: str, page: int) -> list[int]:
    columns = page_columns(pdf_path, page)
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        pdf_page = pdf[page - 1]
        width, _ = pdf_page.get_size()
        textpage = pdf_page.get_textpage()
        text = textpage.get_text_range()
        found = []
        searcher = textpage.search(tag, match_case=True)
        while (hit := searcher.get_next()) is not None:
            start, count = hit
            following = text[start + count : start + count + 1]
            strip = re.fullmatch(r"-X\d+", tag) is not None
            if following.isalnum() or (following in {".", ":"} and not (strip and following == ":")
                                       and text[start + count + 1 : start + count + 2].isalnum()):
                continue
            left = textpage.get_charbox(start)[0] / width
            right = textpage.get_charbox(start + count - 1)[2] / width
            found += [c.n for c in columns if c.x0 <= left <= c.x1 or c.x0 <= right <= c.x1]
        return found
    finally:
        pdf.close()


ALL_REFS = [(code, tag, ref) for code in SETS for tag, ref in _references(code)]


@pytest.mark.parametrize("code", list(SETS))
def test_document_set_is_complete(code):
    missing = [name for name, path in _files(code).items() if not path.exists()]
    assert not missing, f"{code}: fehlt {missing}"
    assert len(_references(code)) > 60


@pytest.mark.parametrize(("code", "tag", "ref"), ALL_REFS, ids=[f"{c} {t} {r}" for c, t, r in ALL_REFS])
def test_reference_points_to_drawn_column(code, tag, ref):
    sheet, column = map(int, REF.match(ref).groups())
    columns = _columns_of(_files(code)["pdf"], tag, sheet)
    assert column in columns, f"{code} {tag} {ref}: im Plan in Spalten {columns}"


@pytest.mark.parametrize("code", list(SETS))
def test_every_symbol_address_is_used_in_the_program(code):
    files = _files(code)
    symbols = parse_symbol_table(read_text(files["sdf"]))
    program = read_text(files["awl"])
    compact = re.sub(r"\s+", "", program)
    for row in symbols:
        address = row["address"]
        if address.startswith(("FB", "DB", "OB", "T ")):
            continue
        assert address.replace(" ", "") in compact, f"{code}: {row['symbol']} {address} nicht im AWL"


@pytest.mark.parametrize("code", list(SETS))
def test_program_parses_with_many_networks(code):
    blocks = parse_awl(read_text(_files(code)["awl"]))
    fb = next(b for b in blocks if b.name.startswith("FB"))
    assert len(fb.networks) >= 10
    assert any(b.name.startswith("OB") for b in blocks)


@pytest.mark.parametrize("code", list(SETS))
def test_signal_path_reaches_the_target_drive(code):
    files = _files(code)
    graph = build_graph(_terminal_rows(files["csv"]), _bom_rows(files["bom"]), parse_symbol_table(read_text(files["sdf"])), read_text(files["awl"]))
    path = signal_path(graph, SETS[code]["start"])
    assert path is not None
    ids = {n["id"] for n in path["nodes"]}
    assert SETS[code]["target"] in ids, f"{code}: {SETS[code]['start']} -> {sorted(ids)[:30]}"
    assert any(n["kind"] == "network" for n in path["nodes"])


@pytest.mark.parametrize("code", list(SETS))
def test_manual_yields_faults_with_tags_and_references(code):
    files = _files(code)
    faults = fault_rows_from_markdown(read_text(files["md"]), "Betriebsanleitung")
    assert len(faults) >= 8
    assert all(f["symptom"] and (f["cause"] or f["fix"]) for f in faults)
    assert sum(1 for f in faults if f["tags"]) >= len(faults) - 1
    assert guess_machine([f"{code} Testdokumentation", "Umroller" if code == "UR-01" else "Aufrollung"])[1] in {"main", "other"}


@pytest.mark.parametrize("code", list(SETS))
def test_terminal_plan_and_bom_only_name_known_devices(code):
    files = _files(code)
    program_tags = {t.tag for t in extract_tags(read_text(files["awl"])) if t.tag_type == TagType.DEVICE}
    bom_tags = {row[0] for row in _bom_rows(files["bom"]) if row[0]}
    for row in _terminal_rows(files["csv"]):
        for cell in row[2:4]:
            for tag in extract_tags(cell):
                if tag.tag_type == TagType.DEVICE and not re.fullmatch(r"-A\d+(\.\d+)?", tag.tag):
                    assert tag.tag in bom_tags, f"{code}: {tag.tag} im Klemmenplan, aber nicht in der Stueckliste"
    assert program_tags <= bom_tags | {"-A1"}, f"{code}: {sorted(program_tags - bom_tags)}"


# --- Befunde aus dem Abschluss-Review ---------------------------------------------------------


def _graph(code):
    files = _files(code)
    return build_graph(_terminal_rows(files["csv"]), _bom_rows(files["bom"]), parse_symbol_table(read_text(files["sdf"])), read_text(files["awl"]))


@pytest.mark.parametrize("code", list(SETS))
def test_safety_device_reaches_the_target_drive(code):
    """Reissleine/Schutztuer -> Sicherheitsrelais -> E-Eingang -> Freigabe -> Umrichter."""
    path = signal_path(_graph(code), SETS[code]["safety"])
    assert path is not None
    ids = {n["id"] for n in path["nodes"]}
    assert SETS[code]["target"] in ids, f"{code}: {SETS[code]['safety']} -> {sorted(ids)[:20]}"


ADDRESS = re.compile(r"([EA]\d+\.\d)")
TERMINAL = re.compile(r"-X3:\d+")


@pytest.mark.parametrize("code", list(SETS))
def test_fault_texts_name_the_terminal_of_the_mentioned_address(code):
    """Jede -X3-Klemme in einer Fehlerzeile gehoert zu einer Adresse, die dieselbe Zeile nennt (oder ist Versorgung)."""
    files = _files(code)
    with files["csv"].open(encoding="utf-8-sig") as handle:
        rows = [row for row in csv.reader(handle, delimiter=";") if len(row) >= 6 and row[1].startswith("-X3:")]
    address_of = {row[1]: m.group(1) for row in rows if (m := ADDRESS.search(row[2]))}
    known = {row[1] for row in rows}
    for fault in fault_rows_from_markdown(read_text(files["md"]), "Betriebsanleitung"):
        text = " ".join((fault["symptom"], fault["cause"], fault["fix"]))
        addresses = set(ADDRESS.findall(text))
        for terminal in set(TERMINAL.findall(text)):
            assert terminal in known, f"{code}: {terminal} gibt es nicht ({fault['symptom']!r})"
            if terminal in address_of:
                assert address_of[terminal] in addresses, f"{code}: {fault['symptom']!r} nennt {terminal}, das ist {address_of[terminal]}"


def test_continuation_sheet_keeps_channel_numbers():
    pdf = pdfium.PdfDocument(str(_files("UR-01")["pdf"]))
    try:
        texts = [pdf[i].get_textpage().get_text_range() for i in range(len(pdf))]
    finally:
        pdf.close()
    page = next(t for t in texts if "E1.4 bis E1.7" in t and "DI 16 x 24 V DC" in t)  # nicht das Inhaltsverzeichnis
    assert re.search(r"13\s*E1\.4", page), "Fortsetzungsblatt muss bei Kanal 13 weiterzaehlen"
    assert "-A1.1" in page


@pytest.mark.parametrize("code", list(SETS))
def test_no_tag_ends_with_a_period(code):
    faults = fault_rows_from_markdown(read_text(_files(code)["md"]), "Betriebsanleitung")
    assert not [t for f in faults for t in f["tags"] if t.endswith(".")]
