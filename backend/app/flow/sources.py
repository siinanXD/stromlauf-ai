"""Dokumente fuer die Extraktion laden: Text je Seite oder Abschnitt, SHA-256 je Datei.

Bewusst ohne Docling: die Extraktion soll in unter 30 s laufen. PDF-Text kommt aus pdfium,
Tabellen aus openpyxl/csv, AWL und Symboltabellen aus den vorhandenen Parsern.
"""

import csv
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from app.ingestion import awl_parser, doctype
from app.models import DocType


@dataclass
class Passage:
    file: str
    page: int | None
    chunk: str  # Abschnittsbezeichnung, z. B. "FB 10 / NW 3" oder "Zeilen 1-40"
    text: str


@dataclass
class DocText:
    file: str
    doc_type: str
    sha256: str
    passages: list[Passage] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.passages)


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def _pdf_passages(path: Path) -> list[Passage]:
    import pypdfium2 as pdfium

    from app.ingestion.docling_parser import pdfium_lock

    with pdfium_lock:
        pdf = pdfium.PdfDocument(str(path))
        try:
            return [
                Passage(path.name, i + 1, f"S. {i + 1}", pdf[i].get_textpage().get_text_range())
                for i in range(len(pdf))
            ]
        finally:
            pdf.close()


def _xlsx_passages(path: Path) -> list[Passage]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines = [
            " | ".join(str(c) for c in row if c is not None)
            for row in workbook.active.iter_rows(values_only=True)
            if any(c is not None for c in row)
        ]
    finally:
        workbook.close()
    return _lines_to_passages(path.name, lines)


def _csv_passages(path: Path) -> list[Passage]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        lines = [" | ".join(row) for row in csv.reader(handle, delimiter=delimiter)]
    return _lines_to_passages(path.name, lines)


def _awl_passages(path: Path) -> list[Passage]:
    passages = []
    for block in awl_parser.parse_awl(awl_parser.read_text(path)):
        if block.declaration:
            passages.append(Passage(path.name, None, block.label, f"Baustein {block.label}\n{block.declaration}"))
        for network in block.networks:
            passages.append(
                Passage(path.name, None, f"{block.label} / NW {network.number}", f"Netzwerk {network.number}: {network.title}\n{network.code}")
            )
    return passages


def _symbol_passages(path: Path) -> list[Passage]:
    rows = awl_parser.parse_symbol_table(awl_parser.read_text(path))
    lines = ["Symbol | Adresse | Datentyp | Kommentar"] + [
        f"{r['symbol']} | {r['address']} | {r['data_type']} | {r['comment']}" for r in rows
    ]
    return _lines_to_passages(path.name, lines, size=60)


def _text_passages(path: Path) -> list[Passage]:
    text = awl_parser.read_text(path)
    if path.suffix.lower() == ".md":
        return _markdown_sections(path.name, text)
    return _lines_to_passages(path.name, text.splitlines())


def _markdown_sections(file: str, text: str) -> list[Passage]:
    passages: list[Passage] = []
    heading, lines = "Anfang", []
    for line in text.splitlines():
        if line.startswith("#"):
            if lines:
                passages.append(Passage(file, None, heading, "\n".join(lines)))
            heading, lines = line.lstrip("# ").strip()[:60], [line]
        else:
            lines.append(line)
    if lines:
        passages.append(Passage(file, None, heading, "\n".join(lines)))
    return passages


def _lines_to_passages(file: str, lines: list[str], size: int = 40) -> list[Passage]:
    return [
        Passage(file, None, f"Zeilen {i + 1}-{min(i + size, len(lines))}", "\n".join(lines[i : i + size]))
        for i in range(0, len(lines), size)
    ]


def load_document(path: Path, doc_type: str | None = None) -> DocText:
    """Text je Seite/Abschnitt; Typ aus Inhalt (doctype.py), wenn nicht vorgegeben."""
    resolved = doc_type or doctype.detect(path.name, path).doc_type
    suffix = path.suffix.lower()
    if resolved == DocType.PLC_PROGRAM or suffix == ".awl":
        passages = _awl_passages(path)
    elif resolved == DocType.PLC_SYMBOLS or suffix == ".sdf":
        passages = _symbol_passages(path)
    elif suffix == ".pdf":
        passages = _pdf_passages(path)
    elif suffix == ".xlsx":
        passages = _xlsx_passages(path)
    elif suffix == ".csv":
        passages = _csv_passages(path)
    else:
        passages = _text_passages(path)
    return DocText(path.name, str(resolved), sha256_of(path), passages)


def render(docs: list[DocText], doc_types: set[str] | None = None, max_chars: int = 120_000) -> str:
    """Kontext fuer das Modell: jede Passage mit Marke [[Datei | Seite/Abschnitt]], damit source belegbar ist."""
    parts: list[str] = []
    total = 0
    for doc in docs:
        if doc_types and doc.doc_type not in doc_types:
            continue
        for passage in doc.passages:
            where = f"S. {passage.page}" if passage.page else passage.chunk
            block = f"[[{passage.file} | {where}]]\n{passage.text.strip()}\n"
            if total + len(block) > max_chars:
                parts.append("[[... gekuerzt ...]]")
                return "\n".join(parts)
            parts.append(block)
            total += len(block)
    return "\n".join(parts)
