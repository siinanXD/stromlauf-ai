"""Dokumenttyp aus dem Inhalt erkennen, nicht aus dem Dateinamen.

Kundendateien heissen "4711_E_Rev3.pdf". Der Typ entscheidet aber, welcher Parser laeuft
(Klemmenplan -> Signalweg, AWL -> Netzwerke) und welche Lueckenregeln der Steckbrief anwendet.
Jede Regel liefert Punkte und einen Grund, der im Upload-Dialog steht ("Kopfzeile Klemmleiste;Klemme").
Bei weniger als MIN_SCORE Punkten bleibt es bei OTHER; dann greift der Dateiname.
"""

import re
from dataclasses import dataclass
from pathlib import Path

from app.models import DocType

MIN_SCORE = 4
SAMPLE_PAGES = 3
SAMPLE_LINES = 300

_FILENAME_HINTS: list[tuple[str, DocType]] = [
    (r"st(ü|ue|u)ckliste|\bbom\b|artikelliste", DocType.BOM),
    (r"klemm", DocType.TERMINAL_PLAN),
    (r"symbol", DocType.PLC_SYMBOLS),
    (r"\bawl\b|\bscl\b|sps|plc", DocType.PLC_PROGRAM),
    (r"stromlauf|schaltplan|eplan|schematic|elektroplan", DocType.SCHEMATIC),
    (r"handbuch|manual|anleitung|betriebsanl|datasheet|datenblatt", DocType.MANUAL),
]

_AWL_BLOCK = re.compile(r"^\s*(FUNCTION_BLOCK|FUNCTION|ORGANIZATION_BLOCK|DATA_BLOCK)\b", re.M)
_AWL_NETWORK = re.compile(r"^\s*NETWORK\b", re.M)
_SYMBOL_ROW = re.compile(
    r'^"[^"\n]*",\s*"\s*(?:%?[EAIQM]\s*\d+\.\d|[EAIQM][BWD]\s*\d+|DB\s*\d+|FB\s*\d+|FC\s*\d+|OB\s*\d+)\s*",\s*"',
    re.M | re.I,
)
_TERMINAL_HEADER = re.compile(r"klemm", re.I)
_TERMINAL_TARGET = re.compile(r"ziel|intern|extern|anschluss|brücke|bruecke", re.I)
_TERMINAL_ROW = re.compile(r"-X\d+:[A-Z0-9]+")
_BOM_TAG = re.compile(
    r"\b(bmk|kennzeichen|betriebsmittel|artikel|pos\.?|label|designation|item)\b", re.I
)
_BOM_QTY = re.compile(
    r"\b(menge|st(ü|ue)ck|stk|anzahl|bestell|hersteller|typ|qty|quantity|manufacturer)\b", re.I
)
_BOM_TITLE = re.compile(
    r"^\W*(artikelst(ü|ue)ckliste|st(ü|ue)ckliste|bauteilliste|nomenclature|bill of materials?|parts?\s*list|artikelliste)",
    re.I,
)
_SHEET_FRAME = re.compile(r"\bblatt\s*\d+\s*/\s*\d+|\bsheet\s*\d+\s*(?:/|of)\s*\d+", re.I)
_FOLIO_FRAME = re.compile(r"\bfolio\s*:\s*\d+", re.I)  # QElectroTech-Schriftfeld
_SHEET_LIST = re.compile(r"\b(folio list|sheet list|inhaltsverzeichnis|table of contents)\b", re.I)
_COLUMN_HEADER = re.compile(
    r"^\s*1\s+2\s+3\s+4\s+5\s+6\s+7\s+8(?:\s+\d{1,2})*\s*$", re.M
)  # Raster 1..8 oder 1..18
_RAILS = re.compile(r"^L1\s*$\s*^L2\s*$\s*^L3\s*$", re.M)
_SCHEMATIC_TITLE = re.compile(
    r"stromlaufplan|schaltplan|circuit diagram|wiring diagram|elektroplan|electrical cabinet", re.I
)
_CROSS_REF = re.compile(r"(?<![\w/.])/\d{1,4}\.\d{1,2}(?![\w.])")
_MANUAL_TITLE = re.compile(
    r"^\W*(betriebsanleitung|bedienungsanleitung|handbuch|manual|operating instructions|datenblatt|datasheet)",
    re.I,
)
_MANUAL_WORDS = re.compile(
    r"\b(wartung|inbetriebnahme|sicherheitshinweis|st(ö|oe)rung|fehlerbehebung|bestimmungsgem|instandhaltung|maintenance|troubleshooting)",
    re.I,
)
_DEVICE_TAG = re.compile(r"(?<![\w.])-[A-Z]{1,3}\d{1,4}(?![\w:])")


@dataclass(frozen=True)
class Detection:
    doc_type: DocType
    confidence: float  # 0..1
    reason: str
    source: str  # content | filename | suffix | none


def sample_text(path: Path) -> str:
    """Kurze Textprobe: erste Seiten/Zeilen. Leer fuer Formate, die nur Docling liest (docx, Bilder)."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pypdfium2 as pdfium

        from app.ingestion.docling_parser import pdfium_lock

        with pdfium_lock:
            pdf = pdfium.PdfDocument(str(path))
            try:
                pages = [
                    pdf[i].get_textpage().get_text_range()
                    for i in range(min(SAMPLE_PAGES, len(pdf)))
                ]
            finally:
                pdf.close()
        return "\n".join(pages)
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines = []
            for row in workbook.active.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
                if len(lines) >= 40:
                    break
            return "\n".join(lines)
        finally:
            workbook.close()
    if suffix in {
        ".csv",
        ".txt",
        ".md",
        ".awl",
        ".scl",
        ".sdf",
        ".log",
        ".ini",
        ".xml",
        ".json",
        ".html",
        ".htm",
    }:
        from app.ingestion.awl_parser import read_text

        return "\n".join(read_text(path).splitlines()[:SAMPLE_LINES])
    return ""


def _score(text: str) -> dict[DocType, tuple[int, list[str]]]:
    lines = [line for line in text.splitlines() if line.strip()]
    head = "\n".join(lines[:5])
    scores: dict[DocType, tuple[int, list[str]]] = {}

    def add(doc_type: DocType, points: int, reason: str) -> None:
        total, reasons = scores.get(doc_type, (0, []))
        scores[doc_type] = (total + points, reasons + [reason])

    if _AWL_BLOCK.search(text):
        add(DocType.PLC_PROGRAM, 10, "AWL-Baustein (FUNCTION_BLOCK/OB/FC)")
    if len(_AWL_NETWORK.findall(text)) >= 2:
        add(DocType.PLC_PROGRAM, 4, "NETWORK-Abschnitte")

    symbol_rows = len(_SYMBOL_ROW.findall(text))
    if symbol_rows >= 3:
        add(DocType.PLC_SYMBOLS, 10, f"{symbol_rows} Symbolzeilen Name, Adresse, Datentyp")

    header = lines[0] if lines else ""
    if _TERMINAL_HEADER.search(header) and _TERMINAL_TARGET.search(header):
        add(DocType.TERMINAL_PLAN, 10, f"Kopfzeile „{header[:60]}“")
    terminal_rows = sum(1 for line in lines if _TERMINAL_ROW.search(line))
    if lines and terminal_rows >= 5 and terminal_rows / len(lines) >= 0.4:
        add(DocType.TERMINAL_PLAN, 6, f"Klemmen -X in {terminal_rows} von {len(lines)} Zeilen")

    for line in lines[:8]:
        if _BOM_TAG.search(line) and _BOM_QTY.search(line):
            add(DocType.BOM, 10, f"Kopfzeile „{line[:60]}“")
            break
    if _BOM_TITLE.search(head):
        add(DocType.BOM, 6, "Titel „Stückliste/Nomenclature“")

    if _SHEET_FRAME.search(text):
        add(DocType.SCHEMATIC, 4, "Schriftfeld „Blatt n / m“")
    if _FOLIO_FRAME.search(text):
        add(DocType.SCHEMATIC, 4, "Schriftfeld „Folio : n“")
    if _SHEET_LIST.search(text):
        add(DocType.SCHEMATIC, 2, "Blattliste/Inhaltsverzeichnis")
    if _COLUMN_HEADER.search(text):
        add(DocType.SCHEMATIC, 4, "Spaltenkopf 1 … 8")
    if _RAILS.search(text):
        add(DocType.SCHEMATIC, 2, "Netzschienen L1 L2 L3")
    if _SCHEMATIC_TITLE.search(head):
        add(DocType.SCHEMATIC, 5, "Titel „Stromlaufplan“")
    refs = len(_CROSS_REF.findall(text))
    if refs >= 3 and DocType.TERMINAL_PLAN not in scores and DocType.BOM not in scores:
        add(DocType.SCHEMATIC, 3, f"{refs} Blattverweise /n.m")

    if _MANUAL_TITLE.search(head):
        add(DocType.MANUAL, 8, "Titel „Betriebsanleitung/Handbuch“")
    words = {m.group(0).lower() for m in _MANUAL_WORDS.finditer(text)}
    if len(words) >= 2:
        add(DocType.MANUAL, 4, "Kapitel zu " + ", ".join(sorted(words)[:3]))
    prose = sum(1 for line in lines if len(line.split()) >= 8)
    if lines and prose / len(lines) >= 0.3 and DocType.PLC_PROGRAM not in scores:
        add(DocType.MANUAL, 3, "Fliesstext")
    return scores


def guess_doc_type(text: str) -> Detection:
    """Bester Typ aus der Textprobe; OTHER, wenn kein Merkmal genug Punkte bringt."""
    if not text.strip():
        return Detection(DocType.OTHER, 0.0, "keine Textprobe", "none")
    scores = _score(text)
    if not scores:
        return Detection(DocType.OTHER, 0.0, "keine bekannten Merkmale", "content")
    ranked = sorted(scores.items(), key=lambda kv: kv[1][0], reverse=True)
    best, (points, reasons) = ranked[0]
    runner_up = ranked[1][1][0] if len(ranked) > 1 else 0
    if points < MIN_SCORE:
        return Detection(DocType.OTHER, 0.0, "zu wenige Merkmale: " + "; ".join(reasons), "content")
    confidence = min(1.0, (points - runner_up) / 10 + 0.3)
    return Detection(best, round(confidence, 2), "; ".join(reasons), "content")


def filename_doc_type(filename: str) -> DocType | None:
    lowered = filename.lower()
    for pattern, doc_type in _FILENAME_HINTS:
        if re.search(pattern, lowered):
            return doc_type
    return None


def detect(filename: str, path: Path | None = None) -> Detection:
    """Endung (.awl/.sdf) > Inhalt > Dateiname > OTHER.

    Ein PDF mit Seiten, aber ohne Text auf den ersten Seiten ist ein Scan: Der Grund sagt das, statt
    "nicht erkannt" (Issue #64); der Typ kommt dann aus dem Dateinamen, sonst OTHER.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".awl":
        return Detection(DocType.PLC_PROGRAM, 1.0, "Endung .awl", "suffix")
    if suffix == ".sdf":
        return Detection(DocType.PLC_SYMBOLS, 1.0, "Endung .sdf", "suffix")
    scan = False
    if path is not None:
        try:
            text = sample_text(path)
            found = guess_doc_type(text)
            scan = suffix == ".pdf" and not text.strip() and _pdf_pages(path) > 0
        except (
            Exception
        ) as exc:  # defekte Datei: Dateiname entscheidet, Ingestion meldet den Fehler
            found = Detection(
                DocType.OTHER, 0.0, f"Textprobe fehlgeschlagen: {type(exc).__name__}", "none"
            )
        if found.doc_type != DocType.OTHER:
            return found
    hinted = filename_doc_type(filename)
    if scan:
        if hinted:
            return Detection(
                hinted, 0.5, f"Scan (keine Textebene); Typ aus Dateiname „{filename}“", "filename"
            )
        return Detection(DocType.OTHER, 0.0, "Scan (keine Textebene)", "content")
    if hinted:
        return Detection(hinted, 0.5, f"Dateiname „{filename}“", "filename")
    return Detection(DocType.OTHER, 0.0, "weder Inhalt noch Dateiname eindeutig", "none")


def _pdf_pages(path: Path) -> int:
    from app.ingestion.docling_parser import pdf_page_count

    return pdf_page_count(path)
