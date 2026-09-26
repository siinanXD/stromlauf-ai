"""Maschine aus der Dokumentation einer Wissensquelle anlegen (Vorschlag + Uebernahme), ohne LLM."""

import re
from pathlib import Path

import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.plant import _get, _machine_out
from app.api.signal import _bom_rows
from app.db import get_session
from app.ingestion.awl_parser import read_text
from app.ingestion.onboarding import fault_rows_from_markdown, guess_machine
from app.ingestion.tags import normalize_tag
from app.models import Chunk, Document, FaultEntry, Hall, KnowledgeSource, Machine, MachineType
from app.schemas import FaultIn, MachineOut

router = APIRouter(prefix="/api", tags=["onboarding"])

TEXT_SUFFIXES = {".md", ".txt", ".markdown"}


class OnboardingDocument(BaseModel):
    filename: str
    doc_type: str


class OnboardingProposal(BaseModel):
    source_id: str
    source_name: str
    name: str
    machine_type: str
    devices: int
    documents: list[OnboardingDocument]
    faults: list[FaultIn]
    hints: list[str]


class OnboardingRequest(BaseModel):
    source_id: str
    name: str
    machine_type: str = "other"
    faults: list[FaultIn] = []


def _label(filename: str) -> str:
    """'06_Betriebsanleitung_FB-01.md' -> 'Betriebsanleitung'."""
    stem = re.sub(r"^\d+[_\- ]+", "", Path(filename).stem)
    return re.split(r"[_ ]", stem, maxsplit=1)[0] or stem


def _markdown_of(session: Session, document: Document) -> str:
    path = Path(document.storage_path)
    if path.suffix.lower() in TEXT_SUFFIXES and path.exists():
        return read_text(path)
    chunks = session.scalars(select(Chunk).where(Chunk.document_id == document.id)).all()
    # Reihenfolge wie im Dokument: Seite, dann Ablagefolge (meta.seq, seit 2026-09 gespeichert)
    chunks = sorted(chunks, key=lambda c: (c.page or 0, int((c.meta or {}).get("seq", 0)), c.id))
    return "\n\n".join(f"## {c.section}\n{c.content}" if c.section else c.content for c in chunks)


def _bom_title(path: Path) -> str:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            first = next(workbook.active.iter_rows(values_only=True), ())
        finally:
            workbook.close()
        return next((str(c).strip() for c in first if c), "")
    except Exception:  # noqa: BLE001 - kaputte Datei: kein Titel
        return ""


@router.get("/sources/{source_id}/onboarding", response_model=OnboardingProposal)
def onboarding_proposal(source_id: str, session: Session = Depends(get_session)):
    source = _get(session, KnowledgeSource, source_id, "Wissensquelle")
    documents = session.scalars(select(Document).where(Document.source_id == source.id).order_by(Document.filename)).all()
    boms = [d for d in documents if d.doc_type == "bom" and Path(d.storage_path).suffix.lower() in {".xlsx", ".xlsm"}]
    bom_title = _bom_title(Path(boms[0].storage_path)) if boms else ""
    devices = 0
    for bom in boms:
        try:
            devices += len(_bom_rows(Path(bom.storage_path)))
        except Exception:  # noqa: BLE001 - kaputte Stueckliste: nicht zaehlen
            continue
    name, machine_type = guess_machine([bom_title, source.name])

    faults: list[dict] = []
    seen: set[str] = set()
    for document in documents:
        if document.doc_type not in {"manual", "other"}:
            continue
        for fault in fault_rows_from_markdown(_markdown_of(session, document), _label(document.filename)):
            if fault["symptom"] not in seen:
                seen.add(fault["symptom"])
                faults.append(fault)

    hints = []
    if not faults:
        hints.append("Keine Fehlertabelle (Symptom | Ursache | Abhilfe) in den Handbüchern gefunden.")
    hints.append("Draufsicht und Schaltschrank-Markierungen können auf der Maschinenseite ergänzt werden (optional per Vision, kostet API-Tokens).")
    return OnboardingProposal(
        source_id=source.id,
        source_name=source.name,
        name=name,
        machine_type=machine_type,
        devices=devices,
        documents=[OnboardingDocument(filename=d.filename, doc_type=d.doc_type) for d in documents],
        faults=[FaultIn(**f) for f in faults],
        hints=hints,
    )


def _unique_name(hall: Hall, name: str) -> str:
    names = {m.name for m in hall.machines}
    if name not in names:
        return name
    number = 2
    while f"{name} ({number})" in names:
        number += 1
    return f"{name} ({number})"


@router.post("/halls/{hall_id}/onboard", response_model=MachineOut, status_code=201)
def onboard(hall_id: str, body: OnboardingRequest, session: Session = Depends(get_session)):
    hall = _get(session, Hall, hall_id, "Halle")
    _get(session, KnowledgeSource, body.source_id, "Wissensquelle")
    if body.machine_type not in set(MachineType):
        raise HTTPException(400, f"machine_type muss eins sein von {sorted(MachineType)}")
    count = len(hall.machines)
    machine = Machine(
        hall_id=hall.id,
        name=_unique_name(hall, body.name.strip() or "Neue Maschine"),
        machine_type=body.machine_type,
        source_id=body.source_id,
        pos_x=72 + (count % 4) * 240,
        pos_y=96 + (count // 4) * 168,
        order_index=count,
        description="Aus der Dokumentation angelegt",
    )
    session.add(machine)
    session.flush()
    for fault in body.faults:
        data = fault.model_dump()
        data["tags"] = [normalize_tag(t) for t in data["tags"] if t.strip()]
        session.add(FaultEntry(machine_id=machine.id, **data))
    session.commit()
    session.refresh(machine)
    return _machine_out(session, machine)
